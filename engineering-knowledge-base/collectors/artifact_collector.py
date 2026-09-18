from __future__ import annotations

import glob
import hashlib
import json
import os
from pathlib import Path
from typing import Any, Iterable

from collectors.base import BaseCollector, KnowledgeDocument


SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".json", ".pdf", ".pptx"}


class ArtifactCollector(BaseCollector):
    """Collect local knowledge artifacts into the existing KnowledgeDocument contract.

    This collector is intentionally acquisition-only. It extracts deterministic text
    and provenance from local files; Architecture Vault enrichment remains responsible
    for candidate generation and Engineering OS remains responsible for canonical
    knowledge promotion.
    """

    def collect(self, source: dict) -> list[KnowledgeDocument]:
        source_name = str(source.get("name") or "").strip()
        if not source_name:
            raise ValueError("artifact source requires a non-empty name")

        paths = self._configured_paths(source)
        files = self._expand_paths(
            paths,
            base_dir=Path(source.get("base_dir") or "."),
            recursive=bool(source.get("recursive", True)),
        )
        if not files:
            raise FileNotFoundError(f"No supported artifact files matched source {source_name!r}")

        tags = [str(item) for item in source.get("tags", []) or []]
        common_metadata = dict(source.get("metadata", {}) or {})
        documents: list[KnowledgeDocument] = []
        for path in files:
            content, extraction_metadata = self._extract(path)
            if not content.strip():
                continue
            raw = path.read_bytes()
            metadata = {
                **common_metadata,
                **extraction_metadata,
                "artifact_path": str(path),
                "artifact_format": path.suffix.lower().lstrip("."),
                "artifact_sha256": hashlib.sha256(raw).hexdigest(),
                "artifact_size_bytes": len(raw),
            }
            documents.append(
                KnowledgeDocument(
                    title=path.stem,
                    url=path.resolve().as_uri(),
                    source_name=source_name,
                    source_type="artifact",
                    content=content,
                    tags=list(tags),
                    metadata=metadata,
                )
            )
        return documents

    @staticmethod
    def _configured_paths(source: dict) -> list[str]:
        raw: Any = source.get("paths")
        if raw is None:
            raw = source.get("path")
        if raw is None:
            raise ValueError("artifact source requires path or paths")
        if isinstance(raw, (str, os.PathLike)):
            return [str(raw)]
        if isinstance(raw, Iterable):
            values = [str(item) for item in raw if str(item).strip()]
            if values:
                return values
        raise ValueError("artifact path/paths must contain at least one path")

    @staticmethod
    def _expand_paths(paths: list[str], *, base_dir: Path, recursive: bool) -> list[Path]:
        resolved_base = Path(os.path.expandvars(os.path.expanduser(str(base_dir)))).resolve()
        discovered: set[Path] = set()
        for configured in paths:
            expanded = os.path.expandvars(os.path.expanduser(configured))
            candidate = Path(expanded)
            if not candidate.is_absolute():
                candidate = resolved_base / candidate

            if any(token in str(candidate) for token in ("*", "?", "[")):
                for match in glob.glob(str(candidate), recursive=True):
                    ArtifactCollector._add_candidate(Path(match), discovered, recursive=recursive)
                continue

            ArtifactCollector._add_candidate(candidate, discovered, recursive=recursive)

        return sorted(
            (path.resolve() for path in discovered if path.suffix.lower() in SUPPORTED_SUFFIXES),
            key=lambda path: str(path).lower(),
        )

    @staticmethod
    def _add_candidate(path: Path, discovered: set[Path], *, recursive: bool) -> None:
        if path.is_file():
            discovered.add(path)
            return
        if not path.is_dir():
            return
        iterator = path.rglob("*") if recursive else path.glob("*")
        for child in iterator:
            if child.is_file() and child.suffix.lower() in SUPPORTED_SUFFIXES:
                discovered.add(child)

    def _extract(self, path: Path) -> tuple[str, dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix in {".md", ".markdown", ".txt"}:
            return path.read_text(encoding="utf-8", errors="ignore"), {}
        if suffix == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            return json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True), {
                "json_top_level_type": type(payload).__name__,
            }
        if suffix == ".pdf":
            return self._extract_pdf(path)
        if suffix == ".pptx":
            return self._extract_pptx(path)
        if suffix == ".ppt":
            raise ValueError("legacy .ppt is not supported; convert to .pptx before ingestion")
        raise ValueError(f"Unsupported artifact type: {suffix}")

    @staticmethod
    def _extract_pdf(path: Path) -> tuple[str, dict[str, Any]]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        sections: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(f"## Page {index}\n\n{text}")
        return "\n\n".join(sections), {"page_count": len(reader.pages)}

    @staticmethod
    def _extract_pptx(path: Path) -> tuple[str, dict[str, Any]]:
        from pptx import Presentation

        presentation = Presentation(str(path))
        slides: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            fragments: list[str] = []
            for shape in slide.shapes:
                text = str(getattr(shape, "text", "") or "").strip()
                if text:
                    fragments.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values = [str(cell.text or "").strip() for cell in row.cells]
                        if any(values):
                            fragments.append(" | ".join(values))
            if fragments:
                slides.append(f"## Slide {index}\n\n" + "\n\n".join(fragments))
        return "\n\n".join(slides), {"slide_count": len(presentation.slides)}
